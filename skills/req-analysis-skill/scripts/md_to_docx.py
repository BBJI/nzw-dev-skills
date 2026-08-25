#!/usr/bin/env python3
"""
nzw-dev-skills Markdown → Docx 转换脚本（req-analysis-skill 用途）

将 PRD.md 及其引用的截图转换为飞书云文档可导入的 .docx 文件。
图片以内联方式嵌入（非外部链接），确保飞书导入后图片完整保留。

依赖：python-docx（缺失时自动安装）
用法（在项目根目录执行）：
  python <skill-dir>/scripts/md_to_docx.py \
    --input .nds/req-NNN/01-requirements/PRD.md \
    --output .nds/req-NNN/01-requirements/PRD.docx
"""

import argparse
import io
import re
import subprocess
import sys
from pathlib import Path

# 修复 Windows 控制台 Unicode 输出
if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

FONT_NAME = 'Microsoft YaHei'
MONO_FONT = 'Consolas'
HEADING_COLOR = (0x1F, 0x23, 0x29)      # 近黑，避免默认主题蓝在飞书中过深
MUTED_COLOR = (0x60, 0x62, 0x66)        # 图注/引用灰
TABLE_HEADER_BG = 'F2F3F5'              # 表头浅灰底纹
RULE_COLOR = 'D0D3D9'                   # 分隔线颜色


def check_dependencies():
    """检查并安装依赖"""
    missing = []
    try:
        from docx import Document  # noqa: F401
    except ImportError:
        missing.append("python-docx")

    if missing:
        print(f"📦 安装缺失依赖: {', '.join(missing)}")
        try:
            subprocess.check_call(
                [sys.executable, "-m", "pip", "install"] + missing,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL
            )
            print("✅ 依赖安装完成")
        except subprocess.CalledProcessError:
            print("❌ 依赖安装失败，请手动安装: pip install " + " ".join(missing))
            sys.exit(1)


def set_run_font(run, name=FONT_NAME, size=None, color=None):
    """统一设置 run 的中西文字体（飞书导入后字体一致）"""
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size is not None:
        run.font.size = size
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def md_to_docx(input_path: str, output_path: str):
    """
    将 Markdown PRD 转换为 .docx

    Args:
        input_path: Markdown 文件路径
        output_path: 输出 .docx 文件路径
    """
    from docx import Document
    from docx.shared import Inches, Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    input_file = Path(input_path)
    base_dir = input_file.parent  # PRD 所在目录，图片相对路径的基准

    if not input_file.exists():
        print(f"❌ 输入文件不存在: {input_file}")
        sys.exit(1)

    md_content = input_file.read_text(encoding="utf-8")
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = FONT_NAME
    font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_NAME)

    # 文档标题属性（飞书导入后显示为文档名）
    title_match = re.search(r'^#\s+(.+)$', md_content, re.MULTILINE)
    if title_match:
        doc.core_properties.title = re.sub(r'`[^`]+`', '', title_match.group(1)).strip()

    lines = md_content.split('\n')
    i = 0
    in_code_block = False
    code_lines = []
    code_lang = ""
    in_table = False
    table_rows = []
    in_blockquote = False
    blockquote_lines = []

    def add_heading(text: str, level: int):
        """添加标题（统一中文字体与深色，飞书映射为标题层级）"""
        clean = re.sub(r'`[^`]+`', lambda m: m.group(0).strip('`'), text)
        clean = clean.strip()
        if clean:
            h = doc.add_heading(clean, level=min(level, 4))
            for run in h.runs:
                set_run_font(run, color=HEADING_COLOR)

    def add_paragraph(text: str, bold: bool = False, italic: bool = False):
        """添加段落（支持 **粗体** / *斜体* / `代码` 行内格式）"""
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
        for part in parts:
            if not part:
                continue
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
                run = p.add_run(part[1:-1])
                run.italic = True
            elif part.startswith('`') and part.endswith('`'):
                run = p.add_run(part[1:-1])
                set_run_font(run, name=MONO_FONT, size=Pt(9.5))
                continue
            else:
                run = p.add_run(part)
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            set_run_font(run)
        return p

    def add_image(image_path: str, alt_text: str = ""):
        """添加图片（内联嵌入 + 居中 + 图注）"""
        img_file = base_dir / image_path
        if img_file.exists():
            try:
                doc.add_picture(str(img_file), width=Inches(6.0))
                last_paragraph = doc.paragraphs[-1]
                last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if alt_text:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    # 注意：对图注 run 单独设置字号，不得修改 Normal 样式本身
                    # （样式级修改会污染全文正文字号）
                    run = cap.add_run(alt_text)
                    set_run_font(run, size=Pt(9), color=MUTED_COLOR)
            except Exception as e:
                doc.add_paragraph(f"[图片: {image_path} (插入失败: {e})]")
        else:
            doc.add_paragraph(f"[图片缺失: {image_path}]")

    def _shade_cell(cell, hex_color: str):
        """给单元格加底纹"""
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement('w:shd')
        shading.set(qn('w:val'), 'clear')
        shading.set(qn('w:fill'), hex_color)
        tc_pr.append(shading)

    def add_table(rows: list):
        """添加表格（Table Grid + 表头加粗底纹）"""
        if not rows:
            return

        # 解析 Markdown 表格行
        parsed_rows = []
        for row in rows:
            cells = [c.strip() for c in row.strip('|').split('|')]
            parsed_rows.append(cells)

        # 过滤分隔行（---）
        data_rows = [r for r in parsed_rows if not all(
            c.replace('-', '').replace(':', '').strip() == '' for c in r
        )]

        if not data_rows:
            return

        # 确定列数并补齐
        max_cols = max(len(r) for r in data_rows)
        for r in data_rows:
            while len(r) < max_cols:
                r.append('')

        table = doc.add_table(rows=len(data_rows), cols=max_cols)
        table.style = 'Table Grid'

        for row_idx, row_data in enumerate(data_rows):
            for col_idx, cell_text in enumerate(row_data):
                if col_idx >= max_cols:
                    continue
                cell = table.cell(row_idx, col_idx)
                # 清理单元格文本中的 Markdown 标记
                clean = cell_text.strip()
                clean = re.sub(r'`([^`]+)`', r'\1', clean)
                clean = re.sub(r'\*\*([^*]+)\*\*', r'\1', clean)
                cell.text = clean
                if row_idx == 0:
                    # 表头：加粗 + 浅灰底纹（飞书导入保留表头样式）
                    _shade_cell(cell, TABLE_HEADER_BG)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True

    def add_horizontal_rule():
        """分隔线渲染为段落底边框（飞书导入为分割线）"""
        p = doc.add_paragraph()
        p_pr = p._p.get_or_add_pPr()
        p_bdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:color'), RULE_COLOR)
        p_bdr.append(bottom)
        p_pr.append(p_bdr)

    def flush_blockquote():
        """刷新引用块"""
        nonlocal in_blockquote, blockquote_lines
        if blockquote_lines:
            text = '\n'.join(blockquote_lines)
            text = re.sub(r'^>\s?', '', text, flags=re.MULTILINE)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            quote_lines = text.split('\n')
            for idx, line in enumerate(quote_lines):
                run = p.add_run(line)
                set_run_font(run, size=Pt(9.5), color=MUTED_COLOR)
                run.italic = True
                if idx < len(quote_lines) - 1:
                    run.add_break()
            blockquote_lines = []
            in_blockquote = False

    def flush_code_block():
        """刷新代码块"""
        nonlocal in_code_block, code_lines, code_lang
        if code_lines:
            if code_lang:
                lang_label = doc.add_paragraph()
                lang_run = lang_label.add_run(f"[{code_lang}]")
                set_run_font(lang_run, size=Pt(8), color=MUTED_COLOR)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            for idx, line in enumerate(code_lines):
                run = p.add_run(line)
                set_run_font(run, name=MONO_FONT, size=Pt(9), color=(0x30, 0x31, 0x33))
                if idx < len(code_lines) - 1:
                    run.add_break()
            code_lines = []
            code_lang = ""
            in_code_block = False

    def flush_table():
        """刷新表格"""
        nonlocal in_table, table_rows
        if table_rows:
            add_table(table_rows)
            table_rows = []
            in_table = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块处理
        if stripped.startswith('```'):
            if in_code_block:
                flush_code_block()
                i += 1
                continue
            else:
                # 先刷新之前的上下文
                flush_blockquote()
                flush_table()
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_lines = []
                i += 1
                continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # 引用块处理
        if stripped.startswith('>'):
            flush_table()
            in_blockquote = True
            blockquote_lines.append(stripped)
            i += 1
            continue

        # 表格行处理
        if stripped.startswith('|') and stripped.endswith('|'):
            flush_blockquote()
            in_table = True
            table_rows.append(stripped)
            i += 1
            continue

        # 非特殊行，先刷新之前的上下文
        flush_blockquote()
        flush_table()

        # 空行
        if not stripped:
            i += 1
            continue

        # 标题
        heading_match = re.match(r'^(#{1,4})\s+(.+)$', stripped)
        if heading_match:
            level = len(heading_match.group(1))
            add_heading(heading_match.group(2), level)
            i += 1
            continue

        # 图片
        img_match = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$', stripped)
        if img_match:
            add_image(img_match.group(2), img_match.group(1))
            i += 1
            continue

        # 分隔线
        if stripped in ('---', '***', '___'):
            add_horizontal_rule()
            i += 1
            continue

        # 任务列表（checkbox），保留缩进层级
        task_match = re.match(r'^(\s*)[-*]\s+\[([ x])\]\s+(.+)$', line)
        if task_match:
            indent = len(task_match.group(1)) // 2
            checked = task_match.group(2) == 'x'
            prefix = "☑" if checked else "☐"
            add_paragraph(f"{'   ' * (indent + 1)}{prefix} {task_match.group(3)}")
            i += 1
            continue

        # 无序列表（含嵌套：按缩进层级递进）
        list_match = re.match(r'^(\s*)[-*]\s+(.+)$', line)
        if list_match:
            indent = len(list_match.group(1)) // 2
            add_paragraph(f"{'   ' * (indent + 1)}• {list_match.group(2)}")
            i += 1
            continue

        # 有序列表
        ol_match = re.match(r'^(\s*)(\d+)\.\s+(.+)$', line)
        if ol_match:
            indent = len(ol_match.group(1)) // 2
            add_paragraph(f"{'   ' * (indent + 1)}{ol_match.group(2)}. {ol_match.group(3)}")
            i += 1
            continue

        # 普通段落
        add_paragraph(stripped)
        i += 1

    # 刷新最后的上下文
    flush_code_block()
    flush_blockquote()
    flush_table()

    # 保存
    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_file))
    print(f"✅ 已生成: {output_file}")
    print(f"📄 飞书导入方式: 飞书云文档 → 新建 → 导入本地文件 → 选择 {output_file.name}")


def main():
    """主函数"""
    parser = argparse.ArgumentParser(description="将 PRD.md 转换为飞书可导入的 .docx（req-analysis-skill）")
    parser.add_argument("--input", required=True, help="输入 Markdown 文件路径")
    parser.add_argument("--output", required=True, help="输出 .docx 文件路径")
    args = parser.parse_args()

    check_dependencies()

    print(f"🔄 转换: {args.input} → {args.output}")
    md_to_docx(args.input, args.output)


if __name__ == "__main__":
    main()
